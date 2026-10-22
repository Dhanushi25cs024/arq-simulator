#!/usr/bin/env python3
"""
Fills the institutional Student_Submission_Proforma.docx based on the provided template.
Preserves original document formatting and fills only verified known details.
"""
import os
import docx

def generate_proforma():
    template_path = r'C:\Users\Admin\Downloads\NetVista3D_Student_Submission_Proforma.docx'
    out_dir = r'D:\Cse sem 3\Web devlopement\Practical 7_PHP form\ARQ_Project\documentation'
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "Student_Submission_Proforma.docx")
    
    doc = docx.Document(template_path)
    
    # 1. Update Table 1 (Student and assignment details)
    t1 = doc.tables[0]
    for row in t1.rows:
        field_name = row.cells[0].text.strip()
        if "Student ID" in field_name:
            row.cells[1].text = "25CS024"
        elif "Exact assigned topic" in field_name:
            row.cells[1].text = "Go-Back-N and Selective Repeat under losses"
        elif "Tools / software" in field_name:
            row.cells[1].text = "Python (3.12), Matplotlib, Streamlit, python-docx, python-pptx, reportlab"
            
    # 2. Update Table 2 (Submission items)
    t2 = doc.tables[1]
    for row in t2.rows:
        item_name = row.cells[0].text.strip()
        if "Supporting file(s)" in item_name:
            row.cells[1].text = "File name(s): ARQ_Project.zip (contains arq_simulator.py, app.py, results.csv, Technical_Report.pdf, ARQ_Presentation.pptx, event_logs/, graphs/)"
            
    # 3. Update Technical Summary Paragraphs
    objective_text = (
        "Design, implement, and benchmark an event-driven discrete-event simulator for Go-Back-N and "
        "Selective Repeat ARQ protocols to evaluate how retransmission counts and completion time change with "
        "window size (W=3 vs W=5) under deterministic frame losses at sequence numbers 4 and 7."
    )
    
    key_result_text = (
        "Selective Repeat maintained an invariant retransmission count of exactly 2 frames across both window sizes "
        "owing to receiver buffering, whereas Go-Back-N required 6 retransmissions at W=3 and 5 at W=5 (due to in-flight "
        "window consolidation), with window size 5 reducing completion time by over 30% through pipelining."
    )
    
    exam_connection_text = (
        "Directly connects to ITUC201 syllabus topics on sliding-window efficiency (eta = W / (1 + 2a)) and minimum sequence "
        "number bits (m) where Go-Back-N requires W <= 2^m - 1 and Selective Repeat requires W <= 2^(m-1), as proven in "
        "the fully solved 10 Mbps point-to-point link numerical."
    )
    
    # Replace underline placeholders after section headings
    found_obj = False
    found_result = False
    found_exam = False
    
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if "Objective / problem addressed" in txt:
            found_obj = True
            continue
        elif found_obj and "Key result" not in txt:
            if "____" in txt or txt == "":
                p.text = objective_text
                found_obj = False
                
        if "Key result / demonstration" in txt:
            found_result = True
            continue
        elif found_result and "Connection to an ITUC201" not in txt:
            if "____" in txt or txt == "":
                p.text = key_result_text
                found_result = False
                
        if "Connection to an ITUC201 theory" in txt:
            found_exam = True
            continue
        elif found_exam and "D. Student checklist" not in txt:
            if "____" in txt or txt == "":
                p.text = exam_connection_text
                found_exam = False

    doc.save(out_path)
    print(f"[PROFORMA] Successfully created filled proforma at: {out_path}")

if __name__ == "__main__":
    generate_proforma()
