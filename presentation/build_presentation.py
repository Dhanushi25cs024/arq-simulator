#!/usr/bin/env python3
"""
Builds the 14-slide academic PowerPoint deck: ARQ_Presentation.pptx
using python-pptx with widescreen layout, formatted tables, and embedded graphs.
"""
import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck(output_pptx: str):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Palette
    NAVY = RGBColor(12, 36, 97)      # #0C2461
    BLUE = RGBColor(30, 55, 153)     # #1E3799
    SLATE = RGBColor(47, 54, 64)     # #2F3640
    ACCENT = RGBColor(39, 174, 96)   # #27AE60 (Green)
    ORANGE = RGBColor(230, 126, 34)  # #E67E22
    LIGHT_BG = RGBColor(248, 249, 250)
    WHITE = RGBColor(255, 255, 255)
    GRAY = RGBColor(113, 128, 147)
    
    blank_layout = prs.slide_layouts[6] # Blank slide
    graphs_dir = os.path.join(os.path.dirname(output_pptx), "..", "graphs")

    def add_header(slide, title_text, category="ITUC201: FUNDAMENTALS OF COMPUTER NETWORKS | STUDENT ID: 25CS024"):
        # Top banner line
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = BLUE
        top_bar.line.fill.background()

        # Category text
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.25), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ORANGE
        
        # Title text
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.7), Inches(0.75))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY

        # Sub-header divider
        divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.02))
        divider.fill.solid()
        divider.fill.fore_color.rgb = RGBColor(220, 221, 225)
        divider.line.fill.background()

    # ----------------------------------------------------------------------------------
    # SLIDE 1: TITLE SLIDE
    # ----------------------------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    # Background accent
    bg_box = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg_box.fill.solid()
    bg_box.fill.fore_color.rgb = LIGHT_BG
    bg_box.line.fill.background()
    
    top_strip = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.2))
    top_strip.fill.solid()
    top_strip.fill.fore_color.rgb = BLUE
    top_strip.line.fill.background()
    
    # Title Card Box
    card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.3), Inches(10.9), Inches(4.9))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = RGBColor(220, 221, 225)
    
    tf1 = card.text_frame
    tf1.word_wrap = True
    p1 = tf1.paragraphs[0]
    p1.text = "NETVISTA 3D: ARQ PROTOCOL EVALUATION"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = NAVY
    p1.alignment = PP_ALIGN.CENTER
    
    p2 = tf1.add_paragraph()
    p2.text = "Comparative Investigation of Go-Back-N and Selective Repeat ARQ under Frame Losses"
    p2.font.size = Pt(15)
    p2.font.color.rgb = BLUE
    p2.alignment = PP_ALIGN.CENTER
    
    p_div = tf1.add_paragraph()
    p_div.text = "__________________________________________________________________"
    p_div.font.size = Pt(11)
    p_div.font.color.rgb = RGBColor(200, 200, 200)
    p_div.alignment = PP_ALIGN.CENTER
    
    p3 = tf1.add_paragraph()
    p3.text = "\nStudent ID: 25CS024  |  Subject: ITUC201 Fundamentals of Computer Networks"
    p3.font.size = Pt(13)
    p3.font.bold = True
    p3.font.color.rgb = SLATE
    p3.alignment = PP_ALIGN.CENTER

    p4 = tf1.add_paragraph()
    p4.text = "Assigned Topic: Go-Back-N & Selective Repeat under losses  |  Assessment: Individual (20 Marks)"
    p4.font.size = Pt(12)
    p4.font.color.rgb = SLATE
    p4.alignment = PP_ALIGN.CENTER

    p5 = tf1.add_paragraph()
    p5.text = "Implementation: Python 3.12 Discrete-Event Engine (Priority Queue)  |  Academic Year: 2026–27"
    p5.font.size = Pt(11)
    p5.font.color.rgb = GRAY
    p5.alignment = PP_ALIGN.CENTER

    # ----------------------------------------------------------------------------------
    # SLIDE 2: OBJECTIVES & PROBLEM STATEMENT
    # ----------------------------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Project Objectives & Core Research Question")
    
    # Left Box: Research Question Callout
    rq_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    rq_box.fill.solid()
    rq_box.fill.fore_color.rgb = RGBColor(235, 240, 255)
    rq_box.line.color.rgb = BLUE
    
    tf_rq = rq_box.text_frame
    tf_rq.word_wrap = True
    p_rq_h = tf_rq.paragraphs[0]
    p_rq_h.text = "PRIMARY RESEARCH QUESTION"
    p_rq_h.font.size = Pt(14)
    p_rq_h.font.bold = True
    p_rq_h.font.color.rgb = BLUE
    
    p_rq_q = tf_rq.add_paragraph()
    p_rq_q.text = "\n\"How do retransmission counts change with window size and loss pattern?\"\n"
    p_rq_q.font.size = Pt(18)
    p_rq_q.font.bold = True
    p_rq_q.font.color.rgb = NAVY
    
    p_rq_desc = tf_rq.add_paragraph()
    p_rq_desc.text = (
        "• Does Go-Back-N's discard policy severely penalize channel efficiency as window size expands?\n\n"
        "• Can Selective Repeat's out-of-order buffering maintain constant retransmission overhead?\n\n"
        "• What happens when multiple frames (4 and 7) are lost in close proximity?"
    )
    p_rq_desc.font.size = Pt(11.5)
    p_rq_desc.font.color.rgb = SLATE

    # Right Box: Objectives
    obj_box = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    obj_box.fill.solid()
    obj_box.fill.fore_color.rgb = WHITE
    obj_box.line.color.rgb = RGBColor(220, 221, 225)
    
    tf_obj = obj_box.text_frame
    tf_obj.word_wrap = True
    p_obj_h = tf_obj.paragraphs[0]
    p_obj_h.text = "TECHNICAL DELIVERABLES & GOALS"
    p_obj_h.font.size = Pt(14)
    p_obj_h.font.bold = True
    p_obj_h.font.color.rgb = NAVY
    
    p_obj_b = tf_obj.add_paragraph()
    p_obj_b.text = (
        "\n1. Event-Driven Simulator (Python 3):\n"
        "   - Pure priority queue (heapq) discrete-event architecture.\n"
        "   - Zero arbitrary loops or polling; explicit simulation timestamps.\n\n"
        "2. Controlled Experimental Framework:\n"
        "   - Window sizes W = 3 vs W = 5.\n"
        "   - Deterministic loss at frames 4 and 7 (first attempt only).\n"
        "   - Identical transmission and propagation latencies.\n\n"
        "3. Verification & Visualizations:\n"
        "   - Space-time ladder diagrams for every experiment.\n"
        "   - Automated CSV metrics, comparative graphs, and Streamlit dashboard.\n\n"
        "4. University Syllabus Grounding:\n"
        "   - Solved ITUC201 examination numerical on link efficiency."
    )
    p_obj_b.font.size = Pt(11)
    p_obj_b.font.color.rgb = SLATE

    # ----------------------------------------------------------------------------------
    # SLIDE 3: ARQ & RELIABLE DATA TRANSMISSION
    # ----------------------------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Reliable Data Transmission & Sliding Window Concept")
    
    # 3 Cards
    cards_data = [
        ("1. Stop-and-Wait ARQ",
         "• Efficiency: η = 1 / (1 + 2a), where a = Tp / Tt.\n\n"
         "• Transmitter waits for ACK after every single frame.\n\n"
         "• If propagation delay (Tp) is high, link remains idle > 90% of the time.\n\n"
         "• Unacceptable for modern high-bandwidth high-latency links.",
         RGBColor(245, 246, 250)),
        ("2. Pipelining & Sliding Window",
         "• Allows up to W frames to be in flight simultaneously.\n\n"
         "• Link Saturation Condition: W ≥ 1 + 2a.\n\n"
         "• Achieves up to 100% link utilization under zero errors.\n\n"
         "• Requires sequence number management and recovery policies under frame loss.",
         RGBColor(235, 240, 255)),
        ("3. Loss Recovery Protocols",
         "• Go-Back-N (GBN):\n"
         "  - Receiver window Wr = 1.\n"
         "  - Discards out-of-order frames.\n"
         "  - Retransmits all in-flight frames.\n\n"
         "• Selective Repeat (SR):\n"
         "  - Receiver window Wr = Ws.\n"
         "  - Buffers out-of-order frames.\n"
         "  - Retransmits solely the lost frame.",
         RGBColor(245, 246, 250))
    ]
    for idx, (head, body, bg_col) in enumerate(cards_data):
        c_shape = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8 + idx*3.95), Inches(1.7), Inches(3.75), Inches(5.1))
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = bg_col
        c_shape.line.color.rgb = RGBColor(220, 221, 225)
        tf = c_shape.text_frame
        tf.word_wrap = True
        p_h = tf.paragraphs[0]
        p_h.text = head
        p_h.font.size = Pt(13)
        p_h.font.bold = True
        p_h.font.color.rgb = NAVY
        p_b = tf.add_paragraph()
        p_b.text = f"\n{body}"
        p_b.font.size = Pt(10.5)
        p_b.font.color.rgb = SLATE

    # ----------------------------------------------------------------------------------
    # SLIDE 4: GO-BACK-N ARQ
    # ----------------------------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Go-Back-N (GBN) Protocol Principles")
    
    gbn_box1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    gbn_box1.fill.solid()
    gbn_box1.fill.fore_color.rgb = WHITE
    gbn_box1.line.color.rgb = BLUE
    tf_g1 = gbn_box1.text_frame
    tf_g1.word_wrap = True
    p = tf_g1.paragraphs[0]
    p.text = "SENDER BEHAVIOR (GBN)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p_b = tf_g1.add_paragraph()
    p_b.text = (
        "\n• Sender Window:\n"
        "  - Window size Ws. Window range: [send_base, send_base + Ws - 1].\n"
        "  - next_seq_num points to next unsent frame.\n\n"
        "• Single Logical Timer:\n"
        "  - Timer is set for the oldest unacknowledged frame (send_base).\n"
        "  - Timer restarts whenever an ACK advances send_base.\n"
        "  - Timer stops only when all sent frames are acknowledged.\n\n"
        "• Retransmission on Timeout:\n"
        "  - If timeout fires for send_base, sender 'Goes Back N'.\n"
        "  - Retransmits ALL outstanding frames currently in flight:\n"
        "    from send_base up to next_seq_num - 1.\n"
        "  - Undamaged frames already sent are wastefully retransmitted!"
    )
    p_b.font.size = Pt(11)
    p_b.font.color.rgb = SLATE

    gbn_box2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    gbn_box2.fill.solid()
    gbn_box2.fill.fore_color.rgb = WHITE
    gbn_box2.line.color.rgb = BLUE
    tf_g2 = gbn_box2.text_frame
    tf_g2.word_wrap = True
    p = tf_g2.paragraphs[0]
    p.text = "RECEIVER BEHAVIOR (GBN)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p_b2 = tf_g2.add_paragraph()
    p_b2.text = (
        "\n• Receiver Window Size Wr = 1:\n"
        "  - Receiver only accepts the exact expected frame: rcv_base.\n"
        "  - Has ZERO buffer capacity for out-of-order packets.\n\n"
        "• In-Order Frame Receipt:\n"
        "  - Accepts frame, passes it to Application Layer.\n"
        "  - Advances rcv_base = rcv_base + 1.\n"
        "  - Sends Cumulative ACK(rcv_base - 1).\n\n"
        "• Out-of-Order Frame Receipt:\n"
        "  - If frame sequence number ≠ rcv_base, frame is DISCARDED.\n"
        "  - Re-transmits duplicate cumulative ACK for the last in-order frame.\n\n"
        "• Sequence Number Requirement:\n"
        "  - Minimum sequence numbers: Ws ≤ 2^m - 1 (m bits required)."
    )
    p_b2.font.size = Pt(11)
    p_b2.font.color.rgb = SLATE

    # ----------------------------------------------------------------------------------
    # SLIDE 5: SELECTIVE REPEAT ARQ
    # ----------------------------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Selective Repeat (SR) Protocol Principles")
    
    sr_box1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    sr_box1.fill.solid()
    sr_box1.fill.fore_color.rgb = WHITE
    sr_box1.line.color.rgb = ACCENT
    tf_s1 = sr_box1.text_frame
    tf_s1.word_wrap = True
    p = tf_s1.paragraphs[0]
    p.text = "SENDER BEHAVIOR (SR)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p_b = tf_s1.add_paragraph()
    p_b.text = (
        "\n• Window Size Ws:\n"
        "  - Permits up to Ws frames in flight.\n\n"
        "• Per-Frame Independent Timers:\n"
        "  - Every transmitted frame has its OWN dedicated timer.\n"
        "  - When ACK(k) arrives, timer for frame k is cancelled.\n\n"
        "• Selective Retransmission:\n"
        "  - When timer for frame k expires, ONLY frame k is retransmitted.\n"
        "  - Unaffected frames in flight or buffered are NOT retransmitted.\n\n"
        "• Window Advancement:\n"
        "  - Window base advances forward ONLY when the oldest frame at send_base is acknowledged."
    )
    p_b.font.size = Pt(11)
    p_b.font.color.rgb = SLATE

    sr_box2 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    sr_box2.fill.solid()
    sr_box2.fill.fore_color.rgb = WHITE
    sr_box2.line.color.rgb = ACCENT
    tf_s2 = sr_box2.text_frame
    tf_s2.word_wrap = True
    p = tf_s2.paragraphs[0]
    p.text = "RECEIVER BEHAVIOR (SR)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p_b2 = tf_s2.add_paragraph()
    p_b2.text = (
        "\n• Receiver Window Size Wr = Ws:\n"
        "  - Can receive out-of-order packets within [rcv_base, rcv_base + Wr - 1].\n\n"
        "• Receiver Buffering:\n"
        "  - Accepts out-of-order frames and stores them in buffer memory.\n"
        "  - Returns an Individual / Selective ACK(k) acknowledging solely frame k.\n\n"
        "• In-Order Application Delivery:\n"
        "  - When missing frame rcv_base arrives, receiver delivers contiguous buffered frames to application.\n"
        "  - Window base rcv_base slides past all delivered frames.\n\n"
        "• Window Constraint (Crucial Rule):\n"
        "  - Ws + Wr ≤ 2^m  ⟹  Ws ≤ 2^(m - 1) (Half of sequence number space)."
    )
    p_b2.font.size = Pt(11)
    p_b2.font.color.rgb = SLATE

    # ----------------------------------------------------------------------------------
    # SLIDE 6: COMPARISON TABLE
    # ----------------------------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Protocol Architectural Matrix: GBN vs Selective Repeat")
    
    rows, cols = 8, 3
    t_shape = s6.shapes.add_table(rows, cols, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.2))
    table = t_shape.table
    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(4.2)
    table.columns[2].width = Inches(4.333)
    
    matrix_data = [
        ("Protocol Feature", "Go-Back-N (GBN)", "Selective Repeat (SR)"),
        ("Receiver Window (Wr)", "1 (Strict in-order)", "Equal to Ws (Out-of-order permitted)"),
        ("Receiver Buffering", "None (Discards out-of-order frames)", "Buffers all valid frames within window"),
        ("Acknowledgment Scheme", "Cumulative (ACK n confirms ≤ n)", "Individual / Selective (ACK n confirms only n)"),
        ("Timer Management", "Single timer for oldest unACKed frame", "Dedicated independent timer per frame"),
        ("Retransmission Action", "Retransmits all outstanding in-flight frames", "Retransmits only the specific lost frame"),
        ("Bandwidth Efficiency", "Lower under loss (wasteful retransmissions)", "High (strictly minimal retransmissions)"),
        ("Sequence Sizing Rule", "Ws ≤ 2^m - 1", "Ws + Wr ≤ 2^m ⟹ Ws ≤ 2^(m - 1)")
    ]
    for r_idx, row in enumerate(matrix_data):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11 if r_idx > 0 else 12)
            if r_idx == 0:
                p.font.bold = True
                p.font.color.rgb = WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
            else:
                p.font.color.rgb = SLATE
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if r_idx % 2 == 1 else RGBColor(245, 246, 250)

    # ----------------------------------------------------------------------------------
    # SLIDE 7: SIMULATOR ARCHITECTURE
    # ----------------------------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Event-Driven Discrete-Event Simulator Architecture")
    
    sim_box1 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    sim_box1.fill.solid()
    sim_box1.fill.fore_color.rgb = WHITE
    sim_box1.line.color.rgb = BLUE
    tf = sim_box1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "EVENT-DRIVEN ENGINE (heapq)"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p_b = tf.add_paragraph()
    p_b.text = (
        "\n• Priority Queue Scheduling:\n"
        "  - Uses Python's heapq priority queue.\n"
        "  - Simulation clock jumps directly to earliest scheduled event.\n"
        "  - Eliminates polling loops, sleep delays, and race conditions.\n\n"
        "• Event Structure:\n"
        "  - (time, event_id, event_type, frame_id, payload)\n"
        "  - event_id acts as strict tie-breaker for simultaneous events.\n\n"
        "• State Machine Integrity:\n"
        "  - Full tracking of in-flight frames, ACK status, buffer slots.\n"
        "  - Sequence validation ensures obsolete timer events are safely ignored."
    )
    p_b.font.size = Pt(11)
    p_b.font.color.rgb = SLATE

    sim_box2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    sim_box2.fill.solid()
    sim_box2.fill.fore_color.rgb = WHITE
    sim_box2.line.color.rgb = BLUE
    tf2 = sim_box2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "SIMULATION EVENT LIFECYCLE"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p_b2 = tf2.add_paragraph()
    p_b2.text = (
        "\n1. FRAME_TX:\n"
        "   - Pushes frame onto channel; schedules FRAME_ARRIVAL.\n"
        "   - If Attempt 1 and frame in {4, 7}, injects loss (no arrival scheduled).\n\n"
        "2. FRAME_ARRIVAL:\n"
        "   - GBN: Check if in-order; accept or discard; schedule ACK.\n"
        "   - SR: Buffer frame; schedule individual ACK; slide rcv_base.\n\n"
        "3. ACK_ARRIVAL:\n"
        "   - Slide send_base forward; manage timers; send new frames.\n\n"
        "4. TIMEOUT:\n"
        "   - GBN: Go-Back-N (retransmit all in-flight frames).\n"
        "   - SR: Selective Retransmission (retransmit single frame)."
    )
    p_b2.font.size = Pt(11)
    p_b2.font.color.rgb = SLATE

    # ----------------------------------------------------------------------------------
    # SLIDE 8: EXPERIMENTAL SETUP
    # ----------------------------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Experimental Setup & Controlled Assumptions")
    
    p_table = s8.shapes.add_table(6, 3, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.1)).table
    p_table.columns[0].width = Inches(3.2)
    p_table.columns[1].width = Inches(2.2)
    p_table.columns[2].width = Inches(6.333)
    
    param_data = [
        ("Parameter", "Configured Value", "Simulation Significance & Notes"),
        ("Total Unique Data Frames", "10 Frames (0 to 9)", "Default assessment payload size; verifies ordered delivery"),
        ("Frame Transmission Time (Tt)", "1.0 Time Unit (TU)", "Time required by transmitter to place full frame onto the wire"),
        ("Propagation Delay (Tp)", "1.5 Time Units (TU)", "One-way signal transit delay across the physical link"),
        ("ACK Transmission Delay", "0.2 TU (Tx) + 1.5 TU (Prop)", "Round-Trip Time RTT ≈ 1.0 + 1.5 + 0.2 + 1.5 = 4.2 TU"),
        ("Retransmission Timeout (RTO)", "6.0 Time Units (TU)", "Configured above RTT (4.2 TU) to prevent spurious timeouts"),
    ]
    for r_idx, row in enumerate(param_data):
        for c_idx, val in enumerate(row):
            cell = p_table.cell(r_idx, c_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11 if r_idx > 0 else 12)
            if r_idx == 0:
                p.font.bold = True
                p.font.color.rgb = WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
            else:
                p.font.color.rgb = SLATE
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if r_idx % 2 == 1 else RGBColor(245, 246, 250)

    # ----------------------------------------------------------------------------------
    # SLIDE 9: LOSS DEMONSTRATION AT FRAMES 4 & 7
    # ----------------------------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Deterministic Loss Injection at Frames 4 and 7")
    
    # Embed Timelines
    t_gbn = os.path.join(graphs_dir, "timeline_GBN_W3.png")
    t_sr = os.path.join(graphs_dir, "timeline_SR_W3.png")
    
    if os.path.exists(t_gbn) and os.path.exists(t_sr):
        s9.shapes.add_picture(t_gbn, Inches(0.8), Inches(1.5), width=Inches(3.8))
        s9.shapes.add_picture(t_sr, Inches(4.8), Inches(1.5), width=Inches(3.8))
        
    info_box = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(1.5), Inches(3.7), Inches(5.3))
    info_box.fill.solid()
    info_box.fill.fore_color.rgb = RGBColor(235, 240, 255)
    info_box.line.color.rgb = BLUE
    tf_info = info_box.text_frame
    tf_info.word_wrap = True
    p = tf_info.paragraphs[0]
    p.text = "LOSS TRACE ANALYSIS"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p_b = tf_info.add_paragraph()
    p_b.text = (
        "\n• Frame 4 First Attempt:\n"
        "  - Dropped by channel (Red X).\n"
        "  - GBN: Receiver discards 5, 6; sends DupACKs.\n"
        "  - SR: Receiver buffers 5, 6; sends selective ACKs.\n\n"
        "• Frame 7 First Attempt:\n"
        "  - Dropped by channel (Red X).\n"
        "  - In GBN W=3: Separate timeout occurs.\n"
        "  - In GBN W=5: Repaired in the same retransmission window!\n\n"
        "• Retransmission Rule:\n"
        "  - Attempt 2 always delivered successfully."
    )
    p_b.font.size = Pt(10.5)
    p_b.font.color.rgb = SLATE

    # ----------------------------------------------------------------------------------
    # SLIDE 10: RESULTS FOR WINDOW SIZE 3
    # ----------------------------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Experimental Results: Window Size W = 3")
    
    w3_table = s10.shapes.add_table(5, 4, Inches(0.8), Inches(1.6), Inches(11.733), Inches(2.6)).table
    w3_data = [
        ("Performance Metric", "Go-Back-N (W=3)", "Selective Repeat (W=3)", "Performance Difference"),
        ("Retransmissions", "6 frames", "2 frames", "SR saves 4 data-frame retransmissions (-66.7%)"),
        ("Total Transmissions", "16 frames", "12 frames", "SR consumes 25% less link bandwidth"),
        ("Timeout Count", "2 timeouts", "2 timeouts", "Equal timeouts (Loss of 4, Loss of 7)"),
        ("Completion Time", "36.8 TU", "27.4 TU", "SR is 25.5% faster in total transfer time")
    ]
    for r_idx, row in enumerate(w3_data):
        for c_idx, val in enumerate(row):
            cell = w3_table.cell(r_idx, c_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            if r_idx == 0:
                p.font.bold = True
                p.font.color.rgb = WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
            else:
                p.font.color.rgb = SLATE
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if r_idx % 2 == 1 else RGBColor(245, 246, 250)

    # Explanation Card below
    expl_card = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.5), Inches(11.733), Inches(2.4))
    expl_card.fill.solid()
    expl_card.fill.fore_color.rgb = WHITE
    expl_card.line.color.rgb = RGBColor(220, 221, 225)
    tf = expl_card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "ANALYSIS FOR WINDOW SIZE 3"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p_b = tf.add_paragraph()
    p_b.text = (
        "• In GBN (W=3), frames 4, 5, 6 were sent in window 1. Loss of 4 caused frames 5 and 6 to be discarded at the receiver. "
        "Upon timeout at t=13.5 TU, GBN retransmitted all 3 frames [4, 5, 6]. Later, frame 7 was lost; timeout triggered retransmission "
        "of frames [7, 8, 9]. Total = 3 + 3 = 6 retransmissions.\n"
        "• In SR (W=3), frames 5 and 6 were buffered. Only frame 4 was retransmitted upon timeout. Later, only frame 7 was retransmitted. "
        "Total retransmissions = 1 + 1 = 2 frames."
    )
    p_b.font.size = Pt(11)
    p_b.font.color.rgb = SLATE

    # ----------------------------------------------------------------------------------
    # SLIDE 11: RESULTS FOR WINDOW SIZE 5
    # ----------------------------------------------------------------------------------
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Experimental Results: Window Size W = 5")
    
    w5_table = s11.shapes.add_table(5, 4, Inches(0.8), Inches(1.6), Inches(11.733), Inches(2.6)).table
    w5_data = [
        ("Performance Metric", "Go-Back-N (W=5)", "Selective Repeat (W=5)", "Performance Difference"),
        ("Retransmissions", "5 frames", "2 frames", "SR saves 3 data-frame retransmissions (-60.0%)"),
        ("Total Transmissions", "15 frames", "12 frames", "SR consumes 20% less link bandwidth"),
        ("Timeout Count", "1 timeout", "2 timeouts", "GBN experienced only 1 timeout (repaired both!)"),
        ("Completion Time", "23.4 TU", "19.0 TU", "SR is 18.8% faster; both much faster than W=3")
    ]
    for r_idx, row in enumerate(w5_data):
        for c_idx, val in enumerate(row):
            cell = w5_table.cell(r_idx, c_idx)
            cell.text = val
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(11)
            if r_idx == 0:
                p.font.bold = True
                p.font.color.rgb = WHITE
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
            else:
                p.font.color.rgb = SLATE
                cell.fill.solid()
                cell.fill.fore_color.rgb = WHITE if r_idx % 2 == 1 else RGBColor(245, 246, 250)

    # Explanation Card below
    expl_card2 = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.5), Inches(11.733), Inches(2.4))
    expl_card2.fill.solid()
    expl_card2.fill.fore_color.rgb = WHITE
    expl_card2.line.color.rgb = RGBColor(220, 221, 225)
    tf = expl_card2.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "KEY DISCOVERY: THE WINDOW CONSOLIDATION EFFECT IN GBN (W=5)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ORANGE
    p_b = tf.add_paragraph()
    p_b.text = (
        "• At W = 5, frames 4, 5, 6, 7, 8 were all transmitted in one continuous pipeline burst.\n"
        "• Frame 4 was dropped, and Frame 7 was also dropped on its first attempt during this burst.\n"
        "• When timeout expired on Frame 4, GBN retransmitted the entire window: [4, 5, 6, 7, 8] (5 frames).\n"
        "• Because Frame 7 was retransmitted in this burst (Attempt 2), it arrived safely! This repaired Frame 7 "
        "simultaneously without ever triggering a second timeout. Thus, GBN experienced only 1 timeout and 5 retransmissions!"
    )
    p_b.font.size = Pt(11)
    p_b.font.color.rgb = SLATE

    # ----------------------------------------------------------------------------------
    # SLIDE 12: COMPARISON GRAPHS
    # ----------------------------------------------------------------------------------
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Visual Comparison & Scientific Interpretation")
    
    c1 = os.path.join(graphs_dir, "comparison_retransmissions.png")
    c4 = os.path.join(graphs_dir, "window_size_sensitivity.png")
    
    if os.path.exists(c1) and os.path.exists(c4):
        s12.shapes.add_picture(c1, Inches(0.8), Inches(1.6), width=Inches(5.7))
        s12.shapes.add_picture(c4, Inches(6.8), Inches(1.6), width=Inches(5.7))
        
    txt_bar = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.6), Inches(11.733), Inches(1.4))
    txt_bar.fill.solid()
    txt_bar.fill.fore_color.rgb = RGBColor(235, 240, 255)
    txt_bar.line.color.rgb = BLUE
    tf = txt_bar.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "RESEARCH QUESTION FINDING SUMMARY"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p_b = tf.add_paragraph()
    p_b.text = (
        "1. Selective Repeat is strictly invariant to window size: 2 retransmissions across both W=3 and W=5.\n"
        "2. Pipelining benefits both protocols: W=5 reduced completion time by 36.4% in GBN and 30.7% in SR.\n"
        "3. Receiver buffering is the decisive differentiator in bandwidth preservation under packet loss."
    )
    p_b.font.size = Pt(10.5)
    p_b.font.color.rgb = SLATE

    # ----------------------------------------------------------------------------------
    # SLIDE 13: SOLVED EXAMINATION NUMERICAL
    # ----------------------------------------------------------------------------------
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "Solved University Examination Numerical (ITUC201 Standard)")
    
    ex_box = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.2))
    ex_box.fill.solid()
    ex_box.fill.fore_color.rgb = WHITE
    ex_box.line.color.rgb = RGBColor(220, 221, 225)
    tf = ex_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "PROBLEM STATEMENT & STEP-BY-STEP SOLUTION"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p_b = tf.add_paragraph()
    p_b.text = (
        "\nProblem: 100 km terrestrial link, 10 Mbps bandwidth, propagation velocity v = 2 × 10^8 m/s, Frame size = 1000 bytes (8000 bits).\n\n"
        "1. Latency Calculations:\n"
        "   - Transmission Delay:  Tt = 8000 bits / (10 × 10^6 bps) = 0.8 ms\n"
        "   - Propagation Delay:   Tp = 100 × 10^3 m / (2 × 10^8 m/s) = 0.5 ms\n"
        "   - Link Parameter a:    a = Tp / Tt = 0.5 ms / 0.8 ms = 0.625\n\n"
        "2. Stop-and-Wait ARQ Efficiency:\n"
        "   - η = 1 / (1 + 2a) = 1 / (1 + 2(0.625)) = 1 / 2.25 = 44.44%\n"
        "   - Achievable Throughput = 0.4444 × 10 Mbps = 4.44 Mbps\n\n"
        "3. Optimal Sliding Window Size for 100% Saturation:\n"
        "   - W_opt ≥ 1 + 2a = 1 + 1.25 = 2.25  ⟹  W_opt = ⌈2.25⌉ = 3 frames\n\n"
        "4. Minimum Sequence Number Bits (m):\n"
        "   - Go-Back-N:       Ws ≤ 2^m - 1  ⟹  3 ≤ 2^m - 1  ⟹  2^m ≥ 4  ⟹  m ≥ 2 bits\n"
        "   - Selective Repeat: 2 Ws ≤ 2^m    ⟹  6 ≤ 2^m      ⟹  2^m ≥ 6  ⟹  m ≥ 3 bits"
    )
    p_b.font.size = Pt(11)
    p_b.font.color.rgb = SLATE

    # ----------------------------------------------------------------------------------
    # SLIDE 14: CONCLUSION & REFERENCES
    # ----------------------------------------------------------------------------------
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "Conclusion, References & Execution Guide")
    
    c_box1 = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    c_box1.fill.solid()
    c_box1.fill.fore_color.rgb = WHITE
    c_box1.line.color.rgb = BLUE
    tf = c_box1.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "RESEARCH CONCLUSIONS"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p_b = tf.add_paragraph()
    p_b.text = (
        "\n• Selective Repeat Superiority under Loss:\n"
        "  Confining retransmission strictly to damaged frames minimizes channel contention and retransmission counts.\n\n"
        "• Window Sizing Trade-offs:\n"
        "  Larger windows improve throughput via pipelining, reducing total completion time across all configurations.\n\n"
        "• Implementation Complexity:\n"
        "  SR requires complex receiver sorting and per-frame timer management, whereas GBN achieves simpler operation at the cost of bandwidth."
    )
    p_b.font.size = Pt(11)
    p_b.font.color.rgb = SLATE

    c_box2 = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    c_box2.fill.solid()
    c_box2.fill.fore_color.rgb = WHITE
    c_box2.line.color.rgb = BLUE
    tf2 = c_box2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "EXECUTION COMMANDS & REFERENCES"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p_b2 = tf2.add_paragraph()
    p_b2.text = (
        "\n• Run Core Simulator (CLI):\n"
        "  python arq_simulator.py\n\n"
        "• Run Interactive Dashboard:\n"
        "  python -m streamlit run app.py\n\n"
        "• Key Academic References:\n"
        "  1. Kurose & Ross, Computer Networking: A Top-Down Approach (8th ed., Pearson).\n"
        "  2. Behrouz A. Forouzan, Data Communications and Networking (5th ed., McGraw-Hill).\n"
        "  3. Tanenbaum & Wetherall, Computer Networks (5th ed., Pearson)."
    )
    p_b2.font.size = Pt(11)
    p_b2.font.color.rgb = SLATE

    # Save
    prs.save(output_pptx)
    print(f"[PPTX] Presentation deck generated successfully at: {output_pptx}")

if __name__ == "__main__":
    out_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ARQ_Presentation.pptx")
    create_deck(out_file)
